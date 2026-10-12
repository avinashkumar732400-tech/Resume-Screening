import os
import json
import time
from pathlib import Path

import streamlit as st
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel
from pypdf import PdfReader
from docx import Document


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="AI Resume Screening System",
    page_icon="🤖",
    layout="wide"
)


# ============================================================
# CUSTOM CSS
# ============================================================

st.markdown("""
<style>

.main {
    padding-top: 1rem;
}

.title {
    text-align: center;
    font-size: 42px;
    font-weight: 700;
    margin-bottom: 5px;
}

.subtitle {
    text-align: center;
    font-size: 18px;
    color: #777;
    margin-bottom: 30px;
}

.card {
    padding: 20px;
    border-radius: 12px;
    border: 1px solid #ddd;
    background-color: #ffffff;
    margin-bottom: 15px;
}

.score {
    font-size: 35px;
    font-weight: bold;
}

.section-title {
    font-size: 24px;
    font-weight: 600;
    margin-top: 20px;
}

</style>
""", unsafe_allow_html=True)


# ============================================================
# HEADER
# ============================================================

st.markdown(
    '<div class="title">🤖 AI Resume Screening System</div>',
    unsafe_allow_html=True
)

st.markdown(
    '<div class="subtitle">AI-powered Job Description & Resume Matching</div>',
    unsafe_allow_html=True
)


# ============================================================
# LOAD ENVIRONMENT
# ============================================================

load_dotenv()

api_key = os.getenv("GROQ_API_KEY")

if not api_key:
    st.error(
        "GROQ_API_KEY not found. Please add GROQ_API_KEY to your .env file."
    )
    st.stop()

client = Groq(api_key=api_key)

model = "openai/gpt-oss-120b"


# ============================================================
# DATA MODELS
# ============================================================

class JobDescription(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float | None
    education_requirements: list[str]
    responsibilities: list[str]


class MatchResult(BaseModel):
    score: float
    details: dict


class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = []


class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None

    total_experience_years: float | None = None

    skills: list[str] = []
    experiences: list[Experience] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []


# ============================================================
# READ PDF
# ============================================================

def read_pdf(file_path):

    reader = PdfReader(file_path)

    text = ""

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text


# ============================================================
# READ DOCX
# ============================================================

def read_docx(file_path):

    document = Document(file_path)

    text = ""

    for paragraph in document.paragraphs:

        if paragraph.text.strip():

            text += paragraph.text + "\n"

    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                if cell.text.strip():

                    text += cell.text + "\n"

    return text


# ============================================================
# READ UPLOADED RESUME
# ============================================================

def read_uploaded_resume(uploaded_file):

    file_name = uploaded_file.name.lower()

    if file_name.endswith(".pdf"):

        reader = PdfReader(uploaded_file)

        text = ""

        for page in reader.pages:

            page_text = page.extract_text()

            if page_text:

                text += page_text + "\n"

        return text

    elif file_name.endswith(".docx"):

        document = Document(uploaded_file)

        text = ""

        for paragraph in document.paragraphs:

            if paragraph.text.strip():

                text += paragraph.text + "\n"

        for table in document.tables:

            for row in table.rows:

                for cell in row.cells:

                    if cell.text.strip():

                        text += cell.text + "\n"

        return text

    return None


# ============================================================
# ANALYZE JOB DESCRIPTION
# ============================================================

def analyze_job_description(job_description):

    job_schema = JobDescription.model_json_schema()

    system_prompt = f"""
You are an expert HR assistant.

Your job is to analyze a job description and extract
structured information from it.

Return ONLY valid JSON matching this schema:

{job_schema}

Rules:

- Do not return the schema itself.
- Do not invent information.
- If minimum experience is not mentioned, return null.
- If information for a list is missing, return an empty list.
"""

    user_prompt = f"""
Analyze the following job description:

{job_description}
"""

    response = client.chat.completions.create(

        model=model,

        messages=[
            {
                "role": "system",
                "content": system_prompt
            },
            {
                "role": "user",
                "content": user_prompt
            }
        ],

        response_format={
            "type": "json_object"
        }
    )

    data = json.loads(
        response.choices[0].message.content
    )

    return JobDescription(**data)


# ============================================================
# PARSE RESUME
# ============================================================

def parse_resume(resume_text):

    resume_schema = Resume.model_json_schema()

    system_prompt = f"""
You are an expert resume parser.

Extract information from the resume based on its meaning.

Return ONLY valid JSON matching this schema:

{resume_schema}

Rules:

1. Do not invent information.
2. If a value is unavailable, return null.
3. If a list has no information, return [].
4. Include internships inside experiences.
5. Extract skills from the entire resume.
"""

    user_prompt = f"""
Parse the following resume:

{resume_text}
"""

    response = client.chat.completions.create(

        model=model,

        messages=[
            {
                "role": "system",
                "content": system_prompt
            },
            {
                "role": "user",
                "content": user_prompt
            }
        ],

        response_format={
            "type": "json_object"
        }
    )

    data = json.loads(
        response.choices[0].message.content
    )

    return Resume(**data)


# ============================================================
# CALCULATE MATCH SCORE
# ============================================================

def final_score(job, resume):

    match_schema = MatchResult.model_json_schema()

    prompt = f"""
You are an HR recruiter.

Compare the candidate's resume with the job description.

JOB DESCRIPTION:

{job.model_dump_json(indent=2)}


CANDIDATE RESUME:

{resume.model_dump_json(indent=2)}


Return JSON matching this schema:

{match_schema}


Give:

1. Candidate name
2. Matching skills
3. Missing important skills
4. Whether experience requirement is met
5. Overall match percentage from 0 to 100
6. A short final verdict

Keep the response concise and easy to read.
"""

    response = client.chat.completions.create(

        model=model,

        messages=[
            {
                "role": "user",
                "content": prompt
            }
        ],

        response_format={
            "type": "json_object"
        }
    )

    data = json.loads(
        response.choices[0].message.content
    )

    return MatchResult(**data)


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:

    st.header("⚙️ Settings")

    st.write("### AI Model")

    st.code(model)

    st.divider()

    st.write("### Supported Files")

    st.write("📄 PDF")

    st.write("📝 DOCX")

    st.divider()

    st.info(
        "Upload resumes and provide a job description "
        "to calculate candidate matching scores."
    )


# ============================================================
# JOB DESCRIPTION SECTION
# ============================================================

st.markdown(
    '<div class="section-title">📋 Job Description</div>',
    unsafe_allow_html=True
)

job_description = st.text_area(

    "Paste the job description here",

    height=300,

    placeholder="Paste the complete job description here..."
)


# ============================================================
# RESUME UPLOAD
# ============================================================

st.markdown(
    '<div class="section-title">📄 Upload Resumes</div>',
    unsafe_allow_html=True
)

uploaded_files = st.file_uploader(

    "Upload PDF or DOCX resumes",

    type=["pdf", "docx"],

    accept_multiple_files=True
)


# ============================================================
# ANALYZE BUTTON
# ============================================================

if st.button(
    "🚀 Analyze Candidates",
    type="primary",
    use_container_width=True
):

    if not job_description.strip():

        st.warning(
            "Please enter a job description."
        )

        st.stop()

    if not uploaded_files:

        st.warning(
            "Please upload at least one resume."
        )

        st.stop()

    # --------------------------------------------------------
    # ANALYZE JOB
    # --------------------------------------------------------

    with st.spinner(
        "Analyzing job description..."
    ):

        try:

            job = analyze_job_description(
                job_description
            )

        except Exception as e:

            st.error(
                f"Job description analysis failed: {e}"
            )

            st.stop()


    # --------------------------------------------------------
    # DISPLAY JOB INFORMATION
    # --------------------------------------------------------

    with st.expander(
        "🔍 Extracted Job Requirements",
        expanded=False
    ):

        col1, col2 = st.columns(2)

        with col1:

            st.write("### Role")

            st.write(job.role)

            st.write("### Required Skills")

            for skill in job.required_skills:

                st.write(f"• {skill}")

            st.write("### Preferred Skills")

            for skill in job.preferred_skills:

                st.write(f"• {skill}")

        with col2:

            st.write("### Minimum Experience")

            if job.minimum_experience is not None:

                st.write(
                    f"{job.minimum_experience} years"
                )

            else:

                st.write("Not specified")

            st.write("### Education")

            for education in job.education_requirements:

                st.write(f"• {education}")


    # --------------------------------------------------------
    # PROCESS RESUMES
    # --------------------------------------------------------

    all_results = []

    progress = st.progress(0)

    status = st.empty()

    total_files = len(uploaded_files)

    for index, uploaded_file in enumerate(uploaded_files):

        status.info(
            f"Processing {uploaded_file.name} "
            f"({index + 1}/{total_files})"
        )

        try:

            # Read resume

            resume_text = read_uploaded_resume(
                uploaded_file
            )

            if not resume_text:

                st.warning(
                    f"Could not read {uploaded_file.name}"
                )

                continue

            # Parse resume

            parsed_resume = parse_resume(
                resume_text
            )

            time.sleep(2)

            # Calculate score

            result = final_score(
                job,
                parsed_resume
            )

            all_results.append({

                "name":
                    parsed_resume.name
                    or uploaded_file.name,

                "score":
                    result.score,

                "details":
                    result.details,

                "file":
                    uploaded_file.name,

                "resume":
                    parsed_resume

            })

        except Exception as e:

            st.error(
                f"Error processing "
                f"{uploaded_file.name}: {e}"
            )

        progress.progress(
            (index + 1) / total_files
        )

    status.success(
        "Resume analysis completed."
    )


    # --------------------------------------------------------
    # SORT RESULTS
    # --------------------------------------------------------

    all_results.sort(
        key=lambda candidate:
        candidate["score"],
        reverse=True
    )


    if not all_results:

        st.error(
            "No resumes were successfully processed."
        )

        st.stop()


    # --------------------------------------------------------
    # SUMMARY METRICS
    # --------------------------------------------------------

    st.markdown(
        '<div class="section-title">📊 Screening Summary</div>',
        unsafe_allow_html=True
    )

    total_candidates = len(all_results)

    average_score = sum(
        candidate["score"]
        for candidate in all_results
    ) / total_candidates

    highest_score = all_results[0]["score"]

    col1, col2, col3 = st.columns(3)

    with col1:

        st.metric(
            "Candidates",
            total_candidates
        )

    with col2:

        st.metric(
            "Average Score",
            f"{average_score:.1f}%"
        )

    with col3:

        st.metric(
            "Highest Score",
            f"{highest_score:.1f}%"
        )


    # --------------------------------------------------------
    # RESULTS TABLE
    # --------------------------------------------------------

    st.markdown(
        '<div class="section-title">🏆 Candidate Ranking</div>',
        unsafe_allow_html=True
    )

    for index, candidate in enumerate(all_results):

        score = candidate["score"]

        if score >= 80:

            status_text = "🟢 Strong Match"

        elif score >= 60:

            status_text = "🟡 Moderate Match"

        else:

            status_text = "🔴 Low Match"

        with st.container(border=True):

            col1, col2, col3 = st.columns(
                [4, 2, 3]
            )

            with col1:

                st.subheader(
                    f"{index + 1}. "
                    f"{candidate['name']}"
                )

                st.caption(
                    candidate["file"]
                )

            with col2:

                st.markdown(
                    f"### {score:.1f}%"
                )

            with col3:

                st.write(status_text)


    # --------------------------------------------------------
    # CANDIDATE DETAILS
    # --------------------------------------------------------

    st.markdown(
        '<div class="section-title">👤 Candidate Details</div>',
        unsafe_allow_html=True
    )

    for candidate in all_results:

        details = candidate["details"]

        with st.expander(
            f"👤 {candidate['name']} "
            f"• {candidate['score']:.1f}%"
        ):

            st.write(
                "**Resume:**",
                candidate["file"]
            )

            if isinstance(details, dict):

                for key, value in details.items():

                    formatted_key = key.replace(
                        "_",
                        " "
                    ).title()

                    st.write(
                        f"**{formatted_key}:**"
                    )

                    if isinstance(value, list):

                        for item in value:

                            st.write(
                                f"• {item}"
                            )

                    elif isinstance(value, dict):

                        st.json(value)

                    else:

                        st.write(value)

            else:

                st.write(details)


    # --------------------------------------------------------
    # TOP 2
    # --------------------------------------------------------

    st.markdown(
        '<div class="section-title">⭐ Top Candidates</div>',
        unsafe_allow_html=True
    )

    top_candidates = all_results[:2]

    for candidate in top_candidates:

        st.success(
            f"**{candidate['name']}** "
            f"→ {candidate['score']:.1f}%"
        )


    # --------------------------------------------------------
    # LOWEST 2
    # --------------------------------------------------------

    if len(all_results) > 2:

        st.markdown(
            '<div class="section-title">📉 Lowest Scores</div>',
            unsafe_allow_html=True
        )

        lowest_candidates = all_results[-2:]

        for candidate in lowest_candidates:

            st.warning(
                f"**{candidate['name']}** "
                f"→ {candidate['score']:.1f}%"
            )